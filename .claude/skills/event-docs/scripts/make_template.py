#!/usr/bin/env python3
"""Turn a completed event Project Plan into a reusable blank template.

A past event makes a better starting point than a blank sheet -- it carries the tab structure,
the header rows, the formatting and the formulas. What it must not carry is the past event's
data, which is what this strips.

What is kept on purpose:
  - Tab structure, headers, column widths, fonts, fills, number formats
  - The task list itself (the tasks ARE the template -- they are not event data)
  - Weeks Out values, and the Due Date formulas that derive from the event date
  - Objectives, which recur near-verbatim across events of a type
  - Notes that state a convention rather than a fact about one event

What is cleared:
  - Event date, venue, budget, KPIs, description
  - Every status (reset to Not Started), every event-specific note
  - All data rows on the tracker tabs, keeping their headers
  - Hyperlinks, which otherwise point at the old event's docs

Two improvements over a raw copy, both deliberate:
  - Due Date formulas get an IF guard, so they sit blank until an event date is entered
    instead of showing a negative-date error in every row of a fresh plan.
  - Budget remaining formulas become SUM over a range instead of naming individual rows
    (=B2-B5-B8 breaks the moment the line items differ from the source event).

Usage:
    python3 make_template.py --source la.xlsx --out proprietary-template.xlsx --type proprietary
    python3 make_template.py --source etail.xlsx --out conference-template.xlsx \
        --type conference --drop-tabs BMW,Rockwood,"The Pit Stop"
"""

import argparse
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

# Notes that describe a standing convention, not a fact about one event. Matched case-insensitively
# as substrings; anything else in a Notes column is treated as event data and cleared.
KEEP_NOTES = [
    "marketing weekly event update",
    "post-event list upload",
    "post-event upload template",
]

# Hints written into the Notes column of the template so it teaches the convention.
NOTE_HINTS = {
    "create sfdc": "Event - <Type> - <M.D.YY> - <Event Name>",
    "create slack channel": "#channel-name, once created",
    "create creative brief": "<Event Name> - Design Brief",
}

# xlsx forbids : \ / ? * [ ] in sheet names, so exporting a Google Sheet silently strips them --
# "1:1 Meetings" arrives as "11 Meetings" and "Shipping/Tracking" as "ShippingTracking". Rename
# them to something legal AND readable, so the template never needs the illegal characters back.
SAFE_TAB_NAMES = {
    "11 Meetings": "1-1 Meetings",
    "ShippingTracking": "Shipping-Tracking",
    "SessionsRountables": "Sessions-Roundtables",
    "PartnershipsSponsors": "Partnerships-Sponsors",
    "SwagActivation": "Swag-Activation",
    "StaffingScheduling": "Staffing-Scheduling",
}

# Portal deadline rows get named after the specific hotel or booth in the source event
# ("JW Cutoff Date", "Omni room block cut off date"). Generalize them so the template reads
# for any show. Matched on the lowercased task text as a substring.
PORTAL_LABEL_FIXES = {
    "jw cutoff": "Hotel room block cut off date",
    "omni room block": "Secondary hotel room block cut off date",
    "tabletop logo": "Tabletop / Booth Logo Deadline",
    "tabletop payment": "Tabletop / Booth Payment",
}

# Header rows per tab: everything below this row is data to clear. A tab absent from this map
# is assumed to have a single header row.
HEADER_ROWS = {
    "Budget Tracker ": 4,
    "Budget Sandbox": 3,
    "Itemized Budget": 2,
    "Final Budget": 2,
    "11 Meetings": 2,
    "1-1 Meetings": 2,
    "Hotel Room Block": 0,   # free-form block; cleared wholesale below
    "ROS": 0,
    "AV ROS": 0,
    "Car Schedule": 0,
}


def clear_cell(cell):
    cell.value = None
    if cell.hyperlink:
        cell.hyperlink = None


def reset_timeline(ws, kind):
    """Reset statuses, rewrite due-date formulas with a guard, and scrub event-specific notes."""
    header_cols = {}
    rows_touched = 0
    for r in range(1, ws.max_row + 1):
        first = ws.cell(row=r, column=1).value
        # Section header rows repeat in conference plans ("Portal Deadlines")
        if first and str(first).strip() in ("Description", "Phase", "Portal Deadlines"):
            header_cols = {
                str(ws.cell(row=r, column=c).value or "").strip().lower(): c
                for c in range(1, ws.max_column + 1)
            }
            continue
        if not first or not header_cols:
            continue

        task = str(first).strip().lower()
        portal = "portal deadlines" in header_cols

        if portal:
            for needle, label in PORTAL_LABEL_FIXES.items():
                if needle in task:
                    ws.cell(row=r, column=1).value = label
                    break

        if (c := header_cols.get("status")):
            ws.cell(row=r, column=c).value = "Not Started"
        if (c := header_cols.get("owner")):
            clear_cell(ws.cell(row=r, column=c))

        if (c := header_cols.get("due date")):
            cell = ws.cell(row=r, column=c)
            wk = header_cols.get("weeks out")
            if portal or not wk:
                # Portal deadlines are set by the show, not by the workback formula.
                cell.value = "TBD - confirm in exhibitor portal" if portal else None
            else:
                col = ws.cell(row=r, column=wk).column_letter
                cell.value = (f"=IF('Project Plan '!$B$1=\"\",\"\","
                              f"'Project Plan '!$B$1-({col}{r}*7))")

        if (c := header_cols.get("notes")):
            cell = ws.cell(row=r, column=c)
            existing = str(cell.value or "").lower()
            hint = next((v for k, v in NOTE_HINTS.items() if k in task), None)
            if hint:
                cell.value = hint
            elif not any(k in existing for k in KEEP_NOTES):
                clear_cell(cell)
        rows_touched += 1
    return rows_touched


def clear_tab(ws, header_rows):
    cleared = 0
    start = 1 if header_rows == 0 else header_rows + 1
    for row in ws.iter_rows(min_row=start):
        for cell in row:
            if cell.value is not None:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    continue
                clear_cell(cell)
                cleared += 1
    return cleared


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--source", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--type", required=True, choices=["proprietary", "conference", "summit"])
    p.add_argument("--drop-tabs", default="", help="Comma-separated tabs to delete entirely.")
    p.add_argument("--keep-objectives", default=None,
                   help="Objectives text for Project Plan B4. Default keeps the source's.")
    args = p.parse_args()

    wb = openpyxl.load_workbook(args.source, data_only=False)

    for name in [t.strip() for t in args.drop_tabs.split(",") if t.strip()]:
        if name in wb.sheetnames:
            del wb[name]
            print(f"  dropped tab {name!r}", file=sys.stderr)
        else:
            sys.exit(f"no tab named {name!r}; tabs are {wb.sheetnames}")

    # --- Project Plan header block: clear the facts, keep the labels and objectives
    pp = wb["Project Plan "]
    for ref in ("B1", "B2", "B3", "B5", "B6"):
        clear_cell(pp[ref])
    # Conference sources append the date range to the label itself ("Event Date: Feb. 23 - 26")
    pp["A1"] = "Event Date"
    if args.keep_objectives:
        pp["B4"] = args.keep_objectives
    # Key Links labels stay as prompts, but must not point at the old event
    for row in pp.iter_rows(min_row=1, max_row=8, min_col=3, max_col=8):
        for cell in row:
            if cell.hyperlink:
                cell.hyperlink = None
            # column E+ on conference plans holds the show URL -- event-specific
            if cell.column >= 5 and cell.value:
                clear_cell(cell)
    print("  reset 'Project Plan ' header block", file=sys.stderr)

    if "Timeline + Tasks" in wb.sheetnames:
        n = reset_timeline(wb["Timeline + Tasks"], args.type)
        print(f"  reset {n} timeline rows", file=sys.stderr)

    # --- Budget tabs: make the remaining-budget formulas range-based instead of row-specific
    if "Budget Tracker " in wb.sheetnames:
        bt = wb["Budget Tracker "]
        clear_cell(bt["A1"])          # held the source event's venue name
        bt["A2"] = "Total Budget"     # source said "[including 10k partner dollars]"
        clear_cell(bt["B2"])          # the source event's budget figure
        bt["C2"] = "Remaining Budget"
        bt["D2"] = "=B2-SUM(B5:B60)"
        bt["E2"] = "Actual Budget Remaining"
        bt["F2"] = "=B2-SUM(D5:D60)"
        bt["A3"], bt["B3"] = "Contingency", 0
        bt["C3"], bt["D3"] = "Contingency Used", 0
        print("  rebuilt 'Budget Tracker ' roll-up formulas as SUM ranges", file=sys.stderr)

    # Conference plans carry three budget tabs whose title and roll-up rows sit above the data,
    # so the generic clear below never reaches them. Their roll-ups also name individual rows
    # (=B2-B4-B5-B7-...), which breaks as soon as the line items differ from the source event.
    if "Budget Sandbox" in wb.sheetnames:
        bs = wb["Budget Sandbox"]
        clear_cell(bs["A1"])                      # "<Event> / <Venue>" title block
        bs["A2"] = "Total Budget"
        clear_cell(bs["B2"])
        bs["C2"] = "Remaining Estimated Budget"
        bs["D2"] = "=B2-SUM(B4:B60)"
        bs["E2"] = "Actual Budget Remaining"
        bs["F2"] = "=B2-SUM(D4:D60)"
        for row in bs.iter_rows(min_row=1, min_col=8, max_col=10):   # sponsorship-offset block
            for cell in row:
                if cell.value is not None and not str(cell.value).startswith("="):
                    clear_cell(cell)
        print("  rebuilt 'Budget Sandbox' roll-up and cleared the offset block", file=sys.stderr)

    for tab, title in (("Itemized Budget", "Budget by Event"), ("Final Budget", "Budget Actuals")):
        if tab in wb.sheetnames:
            wb[tab]["A1"] = title
            print(f"  reset {tab!r} title row", file=sys.stderr)

    # --- Every other tab: clear data, keep headers
    for name in wb.sheetnames:
        if name in ("Project Plan ", "Timeline + Tasks"):
            continue
        n = clear_tab(wb[name], HEADER_ROWS.get(name, 1))
        if n:
            print(f"  cleared {n} values from {name!r}", file=sys.stderr)

    for old, new in SAFE_TAB_NAMES.items():
        if old in wb.sheetnames:
            wb[old].title = new
            print(f"  renamed tab {old!r} -> {new!r}", file=sys.stderr)

    wb.save(args.out)
    print(f"\nwrote {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
