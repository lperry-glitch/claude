#!/usr/bin/env python3
"""Build a blank event Project Plan template from scratch.

Why from scratch rather than stripping a past event (make_template.py): a Google Sheet exported
to xlsx carries ~350-500KB of per-cell styling across its thousand formatted-but-empty rows.
The Drive connector takes upload content only as a base64 string argument, so a large workbook
has to be transcribed by hand into the tool call -- which is where it gets corrupted. A
purpose-built workbook lands around 8KB, small enough to upload reliably, and it comes out
cleaner: no merged cells left over from an old venue name, no stale conditional formats.

Styling matches the source plans: Proxima Nova, the FFF2CC header fill, "mmmm d, yyyy" dates,
"$"#,##0.00 currency, and the real column widths.

The Due Date column is a formula keyed to the event date, exactly as the source plans do it,
but wrapped in an IF so a fresh plan shows blanks instead of a negative-date error in every row:

    =IF('Project Plan '!$B$1="","",'Project Plan '!$B$1-(C2*7))

Enter the event date in 'Project Plan '!B1 and the whole timeline dates itself.

Usage:
    python3 build_blank_template.py --type proprietary --out proprietary-template.xlsx
    python3 build_blank_template.py --type conference  --out conference-template.xlsx
"""

import argparse
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("openpyxl is required:  pip install openpyxl")

sys.path.insert(0, __file__.rsplit("/", 1)[0])
from workback import CONFERENCE, CONFERENCE_PORTAL, PROPRIETARY  # noqa: E402

FONT = "Proxima Nova"
HEADER_FILL = PatternFill("solid", start_color="FFFFF2CC")
DATE_FMT = "mmmm d, yyyy"
MONEY_FMT = '"$"#,##0.00'

OBJECTIVES = {
    "proprietary": ("Strengthening Customer Relationships\nExpand Customer Knowledge of Attentive\n"
                    "Drive New Pipeline\nShowcase Company Culture\nGain Feedback and Insights"),
    "conference": ("Strengthening Customer Relationships\nExpand Customer Knowledge of Attentive\n"
                   "Deepen Pipeline\nShowcase Company Culture\nGain Feedback and Insights"),
}

# Key Links labels, laid out in columns C and D as the source plans do
KEY_LINKS = {
    "proprietary": [("KEY LINKS", "Google Drive"), ("Venue Site", "Splash Page"),
                    ("Activation Site", "Invite List")],
    "conference": [("KEY LINKS", "Google Drive"), ("Event Portal", "Event Website"),
                   ("Alliance Portal", "Invite List | Meeting Tracker | RSVPs")],
}

# tab name -> (headers, column widths)
TABS = {
    "proprietary": [
        ("Budget Tracker ", ["Item", "Estimated Cost", "Description", "Actual Cost",
                             "Payment Status", "Notes"], [28, 16, 34, 14, 16, 30]),
        ("Agenda", ["Time", "Description", "Owner"], [18, 42, 18]),
        ("Sessions", ["Session Title", "Abstract", "Speaker", "Speaker Title", "Time"],
         [28, 44, 20, 26, 20]),
        ("1-1 Meetings", ["Company", "Name", "Title", "Time requested", "Time Assigned",
                          "CS Assigned"], [24, 20, 26, 15, 15, 15]),
        ("Printed Assets", ["Asset", "Description", "Quantity", "Payment Status"], [28, 38, 11, 16]),
        ("Digital Assets", ["Item", "Use Case"], [28, 38]),
        ("Shipping-Tracking", ["Item", "Qty", "Tracking Link", "Status"], [28, 8, 30, 14]),
    ],
    "conference": [
        ("Budget Sandbox", ["Item", "Estimated Cost", "Description", "Actual Cost",
                            "Payment Status", "Notes"], [30, 16, 30, 14, 16, 28]),
        ("Itemized Budget", ["Item", "Estimated", "Actual", "Budget Owner", "Notes"],
         [34, 14, 14, 16, 30]),
        ("Final Budget", ["Item", "Estimated", "Actual", "Quoted", "Status Notes",
                          "Invoicing Notes", "Payment Status"], [34, 13, 13, 13, 22, 22, 16]),
        ("Staffing & Passes", ["Pass Type", "Name", "Room Block", "Lodging", "Arrival",
                               "Departure"], [18, 24, 12, 16, 10, 11]),
        ("Speakers", ["Name", "Company", "Title", "Travel Notes", "Headshot / Bio",
                      "Speaker Portal"], [22, 20, 24, 24, 26, 16]),
        ("Sponsors", ["Opportunity", "Partner", "Contract", "Logo", "Splash Page"],
         [26, 20, 14, 10, 14]),
        ("Hotel Room Block", ["Guest Name", "Email", "Confirmation #", "Arrival", "Departure",
                              "Notes"], [22, 30, 16, 10, 11, 24]),
        ("Printed Assets", ["Description", "QTY", "Size / Spec", "Tracking Number", "Cost",
                            "Delivery Status", "Notes"], [28, 8, 22, 22, 10, 16, 24]),
        ("Shipping-Tracking", ["Item", "Qty", "Tracking Information", "Status"], [28, 8, 30, 14]),
        ("ROS", ["Time", "Duration", "Segment", "Owner", "Details/Notes"], [14, 12, 28, 16, 44]),
        ("Vendor List", ["Vendor", "Details", "Contact", "Timings"], [24, 24, 32, 30]),
        ("Car Schedule", ["Time", "# of Cars", "Notes", "Total Hours Used"], [20, 11, 44, 16]),
    ],
}


def style_header(ws, headers, widths):
    for i, (text, width) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=text)
        c.font = Font(name=FONT, bold=True)
        c.fill = HEADER_FILL
        ws.column_dimensions[c.column_letter].width = width


def build_project_plan(wb, kind):
    ws = wb.active
    ws.title = "Project Plan "
    for i, label in enumerate(["Event Date", "Venue", "Budget", "Objectives", "KPIs",
                               "Description:"], start=1):
        c = ws.cell(row=i, column=1, value=label)
        c.font = Font(name=FONT, bold=True)
        c.fill = HEADER_FILL
        ws.cell(row=i, column=2).font = Font(name=FONT)
    ws["B1"].number_format = DATE_FMT
    ws["B1"].alignment = Alignment(horizontal="center")
    ws["B3"].number_format = MONEY_FMT
    ws["B4"] = OBJECTIVES[kind]
    for ref in ("B4", "B5", "B6"):
        ws[ref].alignment = Alignment(wrap_text=True, vertical="top")

    for i, (left, right) in enumerate(KEY_LINKS[kind], start=1):
        for col, text in ((3, left), (4, right)):
            c = ws.cell(row=i, column=col, value=text)
            c.font = Font(name=FONT, bold=(i == 1))
            c.fill = HEADER_FILL
    for col, width in zip("ABCD", (45.5, 67.9, 17.2, 30.0)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[4].height = 76
    ws.row_dimensions[6].height = 60


def build_timeline(wb, kind):
    ws = wb.create_sheet("Timeline + Tasks")
    style_header(ws, ["Description", "Status", "Weeks Out", "Due Date", "Notes"],
                 [40.4, 14.6, 14.0, 19.8, 62.8])
    tasks = sorted(PROPRIETARY if kind == "proprietary" else CONFERENCE, key=lambda r: -r[0])
    row = 2
    for weeks, task, notes in tasks:
        write_task_row(ws, row, task, weeks, notes)
        row += 1

    if kind == "conference":
        row += 1
        for i, text in enumerate(["Portal Deadlines", "Status", "Weeks Out", "Due Date", "Notes"],
                                 start=1):
            c = ws.cell(row=row, column=i, value=text)
            c.font = Font(name=FONT, bold=True)
            c.fill = HEADER_FILL
        row += 1
        for weeks, task in CONFERENCE_PORTAL:
            # The show sets these dates, not the workback formula.
            write_task_row(ws, row, task, weeks, "", due="TBD - confirm in exhibitor portal")
            row += 1


def write_task_row(ws, row, task, weeks, notes, due=None):
    ws.cell(row=row, column=1, value=task).font = Font(name=FONT)
    ws.cell(row=row, column=2, value="Not Started").font = Font(name=FONT)
    ws.cell(row=row, column=3, value=weeks).font = Font(name=FONT)
    c = ws.cell(row=row, column=4,
                value=due or f"=IF('Project Plan '!$B$1=\"\",\"\","
                             f"'Project Plan '!$B$1-(C{row}*7))")
    c.font = Font(name=FONT)
    if not due:
        c.number_format = DATE_FMT
    if notes:
        ws.cell(row=row, column=5, value=notes).font = Font(name=FONT)


def build_budget_rollup(wb, kind):
    """Insert the roll-up rows above the item header on the primary budget tab."""
    name = "Budget Tracker " if kind == "proprietary" else "Budget Sandbox"
    ws = wb[name]
    ws.insert_rows(1, 3)
    pairs = [("Total Budget", None, "Remaining Budget", "=B2-SUM(B5:B60)",
              "Actual Budget Remaining", "=B2-SUM(D5:D60)"),
             ("Contingency", 0, "Contingency Used", 0, None, None)]
    for r, vals in enumerate(pairs, start=2):
        for i, v in enumerate(vals, start=1):
            if v is None:
                continue
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, bold=(i % 2 == 1))
            if i % 2 == 1:
                c.fill = HEADER_FILL
            else:
                c.number_format = MONEY_FMT
    # Pre-format a modest block of item rows as currency; the planner can drag it down.
    # Keeping this small matters: every styled cell is bytes in the upload payload, which the
    # Drive connector accepts only as an inline base64 argument.
    for r in range(5, 25):
        ws.cell(row=r, column=2).number_format = MONEY_FMT
        ws.cell(row=r, column=4).number_format = MONEY_FMT


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--type", required=True, choices=["proprietary", "conference"])
    p.add_argument("--out", required=True)
    args = p.parse_args()

    wb = openpyxl.Workbook()
    build_project_plan(wb, args.type)
    build_timeline(wb, args.type)
    for name, headers, widths in TABS[args.type]:
        style_header(wb.create_sheet(name), headers, widths)
    build_budget_rollup(wb, args.type)

    wb.save(args.out)
    print(f"wrote {args.out}: {len(wb.sheetnames)} tabs -> {wb.sheetnames}", file=sys.stderr)


if __name__ == "__main__":
    main()
